"""Reference solutions for the eval self-test.

Each case carries a ``ref`` spec; applying it must make every check of the
case pass. This proves the checks (and the case paths) are correct before
the agent is graded against them. Runs inside the agent image.
"""
import os
import shutil
import subprocess
import sys
from pathlib import Path

from docgen import docx as D, one, pdf as F, pptx as P, theme as T, xlsx as X  # noqa: E402
from docgen import xlsx_struct as XS  # noqa: E402
from docgen.package import Package  # noqa: E402

SCRIPTS = Path(os.environ.get("DOCGEN_SCRIPTS", "/opt/docgen/scripts"))


def out_name(src, suffix=None):
    p = Path(src)
    return f"{p.stem}-duzenlenmis{suffix or p.suffix}"


# ---- DOCX -----------------------------------------------------------------
def _docx_op(pkg, op):
    k = op["op"]
    if k == "run_props":
        ms = D.find(pkg, op["text"], section=op.get("section"), whole_word=op.get("whole_word", False))
        if op.get("near"):
            ms = [m for m in ms if op["near"] in m.para.text]
        for m in (ms if op.get("all") else [one(ms, op["text"])]):
            D.set_run_props(D.split_runs(m), pkg=pkg, **op["props"])
    elif k == "para_runs":
        para = one(D.find_paragraphs(pkg, op["contains"]), op["contains"])
        D.set_run_props([r.el for r in para.runs], pkg=pkg, **op["props"])
    elif k == "para_props":
        para = one(D.find_paragraphs(pkg, op["contains"], headings_only=op.get("headings_only", False)), op["contains"])
        D.set_paragraph_props([para], **op["props"])
    elif k == "style_levels":
        for lvl in op["levels"]:
            for sid in D.style_ids_for_level(pkg, lvl):
                D.set_style_props(pkg, sid, **op["props"])
    elif k == "style":
        D.set_style_props(pkg, op["style"], **op["props"])
    elif k == "replace":
        D.replace(one(D.find(pkg, op["old"], section=op.get("section"), whole_word=op.get("whole_word", False),
                             part=op.get("part")), op["old"]), op["new"])
    elif k == "replace_all":
        D.replace_all(pkg, op["old"], op["new"], whole_word=op.get("whole_word", True),
                      headings_only=op.get("headings_only", False))
    elif k == "table_cell":
        t = one(D.find_table(pkg, op["table_has"]))
        rows = D.table_rows(t)
        D.set_cell_text(t, [r[0] for r in rows].index(op["row"]),
                        next(i for i, h in enumerate(rows[0]) if op["col"] in h), op["value"])
    elif k == "table_add_row":
        t = one(D.find_table(pkg, op["table_has"]))
        rows = D.table_rows(t)
        D.add_row(t, [r[0] for r in rows].index(op["after"]), op["values"])
    elif k == "table_add_col":
        t = one(D.find_table(pkg, op["table_has"]))
        rows = D.table_rows(t)
        D.add_column(t, len(rows[0]) - 1, [op["header"]] + [""] * (len(rows) - 1))
    elif k == "table_del_col":
        t = one(D.find_table(pkg, op["table_has"]))
        D.delete_column(t, next(i for i, h in enumerate(D.table_rows(t)[0]) if op["col"] in h))
    elif k == "table_header_fill":
        t = one(D.find_table(pkg, op["table_has"]))
        for c in range(len(D.table_rows(t)[0])):
            D.set_cell_shading(t, 0, c, op["color"])
    elif k == "delete_section":
        h = one([x for x in D.headings(pkg) if x.text.strip() == op["heading"]], op["heading"])
        D.delete_section(pkg, h)
    elif k == "append_section":
        h = one([x for x in D.headings(pkg) if x.text.strip() == op["heading"]], op["heading"])
        like = next(b for b in reversed(D.section_blocks(pkg, h)) if b.tag == D.w("p") and "".join(b.itertext()).strip())
        D.append_to_section(pkg, h, D.new_paragraph(op["text"], like=like if like is not h.el else None))
    elif k == "set_para":
        h = one([x for x in D.headings(pkg) if x.text.strip() == op["heading"]], op["heading"])
        D.set_paragraph_text(D.section_blocks(pkg, h)[1], op["text"])
    elif k == "rename_heading":
        D.set_paragraph_text(one([x for x in D.headings(pkg) if x.text.strip() == op["old"]], op["old"]), op["new"])
    elif k == "theme_color":
        T.set_color(pkg, op["slot"], op["color"])
    elif k == "theme_fonts":
        T.set_fonts(pkg, major=op.get("major"), minor=op.get("minor"))
    else:
        raise ValueError(k)


def docx_ops(up, out, file, ops):
    pkg = Package(up / file)
    for op in ops:
        _docx_op(pkg, op)
    pkg.save(out / out_name(file, ".docx"))


def docx_template(up, out, file, title, sections, name="sablon-belge.docx"):
    import docx as pydocx
    dst = out / name
    D.start_from_template(str(up / file), str(dst))
    d = pydocx.Document(str(dst))
    d.add_paragraph(title, style="Title")
    for s in sections:
        d.add_paragraph(s, style="Heading 1")
        d.add_paragraph(f"{s} bölümüne ilişkin açıklama.", style="Normal")
    d.save(str(dst))


def docx_create(up, out, name, title, sections, heading_color=None, table=False, numbered=False):
    import docx as pydocx
    from docx.shared import RGBColor
    d = pydocx.Document()
    d.add_heading(title, 0)
    for s in sections:
        h = d.add_heading(s, 1)
        if heading_color:
            for r in h.runs:
                r.font.color.rgb = RGBColor.from_string(heading_color)
        d.add_paragraph(f"{s} hakkında kısa açıklama.")
    if table:
        t = d.add_table(rows=3, cols=2)
        for i, (a_, b_) in enumerate([("Bölge", "Ciro"), ("Marmara", "5.200.000"), ("Ege", "3.100.000")]):
            t.cell(i, 0).text, t.cell(i, 1).text = a_, b_
    if numbered:
        for i in range(3):
            d.add_paragraph(f"Madde {i + 1}", style="List Number")
    d.save(str(out / name))


# ---- PPTX -----------------------------------------------------------------
def _slide_index(pkg, op):
    return op["slide"] if "slide" in op else P.slide(pkg, title=op["title"]).index


def _shape(pkg, op):
    idx = _slide_index(pkg, op)
    if op.get("kind"):
        return one(P.find_shapes(pkg, slide=idx, kind=op["kind"]))
    if op.get("placeholder"):
        return one(P.find_shapes(pkg, slide=idx, placeholder=op["placeholder"]))
    return one(P.find_shapes(pkg, slide=idx, text=op["shape_text"]))


def _pptx_op(pkg, op):
    k = op["op"]
    if k == "replace":
        P.replace(one(P.find(pkg, op["old"], slide=_slide_index(pkg, op), in_title=op.get("in_title"))), op["new"])
    elif k == "run_props":
        if op.get("text"):
            ms = P.find(pkg, op["text"], slide=_slide_index(pkg, op), in_title=op.get("in_title"),
                        whole_word=op.get("whole_word", False))
            P.set_run_props(P.split_runs(one(ms)), **op["props"])
        else:
            P.set_run_props(P.runs_of(_shape(pkg, op)), **op["props"])
    elif k == "align":
        P.set_paragraph_props(P.paragraphs_of(_shape(pkg, op)), align=op["align"])
    elif k == "move":
        sh = _shape(pkg, op)
        W, H = P.slide_size(pkg)
        x, y, w, h = sh.xfrm
        m = P.cm(1)
        pos = {"top-right": (W - w - m, m), "top-left": (m, m), "bottom-left": (m, H - h - m),
               "bottom-right": (W - w - m, H - h - m), "center": ((W - w) // 2, (H - h) // 2)}[op["pos"]]
        P.set_geometry(sh, x=pos[0], y=pos[1])
    elif k == "widen":
        sh = _shape(pkg, op)
        P.set_geometry(sh, w=int(sh.xfrm[2] * op["factor"]))
    elif k == "scale":
        P.scale(_shape(pkg, op), op["factor"])
    elif k == "fill":
        P.set_fill(_shape(pkg, op), op["color"])
    elif k == "bg":
        P.set_background(P.slide(pkg, title=op["title"]), op["color"])
    elif k == "notes":
        P.set_notes(pkg, op["slide"], op["text"])
    elif k == "table_cell":
        t = one(P.tables(pkg, slide=op["slide"]))
        rows = P.table_rows(t)
        P.set_cell_text(t, [r[0] for r in rows].index(op["row"]), rows[0].index(op["col"]), op["value"])
    elif k == "table_header":
        t = one(P.tables(pkg, slide=op["slide"]))
        for c in range(len(P.table_rows(t)[0])):
            tc = P.cell(t, 0, c)
            if op.get("bold"):
                P.set_run_props(P.runs_of(tc), bold=True)
            if op.get("align"):
                P.set_paragraph_props(P.paragraphs_of(tc), align=op["align"])
    elif k == "table_add_col":
        t = one(P.tables(pkg, slide=op["slide"]))
        rows = P.table_rows(t)
        P.add_column(t, len(rows[0]) - 1, [op["header"]] + [""] * (len(rows) - 1))
    elif k == "table_cell_fill":
        t = one(P.tables(pkg, slide=op["slide"]))
        rows = P.table_rows(t)
        P.set_cell_fill(t, [r[0] for r in rows].index(op["row"]), rows[0].index(op["col"]), op["color"])
    elif k == "delete_slide":
        P.delete_slide(pkg, P.slide(pkg, title=op["title"]).index)
    elif k == "move_slide":
        P.move_slide(pkg, P.slide(pkg, title=op["title"]).index, P.slide(pkg, title=op["before"]).index)
    elif k == "duplicate":
        P.duplicate_slide(pkg, P.slide(pkg, title=op["title"]).index)
    elif k == "hide":
        P.set_hidden(pkg, P.slide(pkg, title=op["title"]).index, op.get("hidden", True))
    elif k == "theme_color":
        T.set_color(pkg, op["slot"], op["color"])
    elif k == "theme_fonts":
        T.set_fonts(pkg, major=op.get("major"), minor=op.get("minor"))
    else:
        raise ValueError(k)


def pptx_ops(up, out, file, ops):
    pkg = Package(up / file)
    for op in ops:
        _pptx_op(pkg, op)
    pkg.save(out / out_name(file, ".pptx"))


def pptx_template(up, out, file, titles, name="sablon-sunum.pptx"):
    from pptx import Presentation
    dst = out / name
    P.start_from_template(str(up / file), str(dst))
    prs = Presentation(str(dst))
    for i, t in enumerate(titles):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = t
        if i:
            s.placeholders[1].text_frame.text = f"{t} için ana mesaj"
    prs.save(str(dst))


def pptx_create(up, out, name, titles, chart=False):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = Presentation()
    for i, t in enumerate(titles):
        s = prs.slides.add_slide(prs.slide_layouts[0 if i == 0 else 1])
        s.shapes.title.text = t
        if i:
            s.placeholders[1].text_frame.text = "Ana mesaj"
    if chart:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Grafik"
        cd = CategoryChartData()
        cd.categories = ["Q1", "Q2", "Q3"]
        cd.add_series("Ciro", (10, 12, 14))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
    prs.save(str(out / name))


# ---- XLSX -----------------------------------------------------------------
def _recalc(path):
    subprocess.run([sys.executable, str(SCRIPTS / "xlsx_recalc.py"), str(path)], check=True, capture_output=True)


def xlsx_ops(up, out, file, ops, recalc=True, onay=False):
    pkg = Package(up / file)
    for op in ops:
        k = op["op"]
        if k == "set_value":
            X.set_value(pkg, op["cell"], op["value"], allow_formula_overwrite=onay)
        elif k == "set_formula":
            X.set_formula(pkg, op["cell"], op["formula"], allow_formula_overwrite=onay)
        elif k == "style":
            X.set_style(pkg, op["range"], font=op.get("font"), fill=op.get("fill"), align=op.get("align"),
                        number_format=op.get("numfmt"))
        elif k == "cf":
            X.add_conditional_format(pkg, op["range"], operator=op.get("operator", "greaterThan"),
                                     formula=op.get("formula"), fill=op.get("fill"), font_color=op.get("font_color"),
                                     expression=op.get("expression"))
        elif k == "width":
            X.set_column_width(pkg, op["sheet"], op["col"], op["width"])
        elif k == "insert_rows":
            XS.insert_rows(pkg, op["sheet"], op["at"], op.get("n", 1))
        elif k == "delete_rows":
            XS.delete_rows(pkg, op["sheet"], op["at"], op.get("n", 1), force=onay)
        elif k == "insert_cols":
            XS.insert_cols(pkg, op["sheet"], op["at"], op.get("n", 1))
        elif k == "sort":
            XS.sort_range(pkg, op["sheet"], op["range"], op["by"], descending=op.get("desc", False))
        else:
            raise ValueError(k)
    dst = out / out_name(file, ".xlsx")
    pkg.save(dst)
    if recalc:
        _recalc(dst)


def xlsx_template(up, out, file, values, name="teklif.xlsx"):
    pkg = Package(up / file)
    if pkg.is_template():
        pkg.as_document()
    for cell, v in values.items():
        X.set_value(pkg, cell, v)
    pkg.save(out / name)
    _recalc(out / name)


def xlsx_create(up, out, name, kind):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    if kind == "budget":
        ws.title = "Bütçe"
        ws.append(["Kalem", "Q1", "Q2", "Q3", "Q4", "Toplam"])
        for i, k in enumerate(["Personel", "Pazarlama", "Ofis", "Yazılım"], start=2):
            ws.append([k, 100 * i, 110 * i, 120 * i, 130 * i, f"=SUM(B{i}:E{i})"])
        ws.append(["Toplam"] + [f"=SUM({c}2:{c}5)" for c in "BCDEF"])
    elif kind == "sales":
        ws.title = "Satışlar"
        ws.append(["Tarih", "Ürün", "Adet", "Fiyat", "Tutar"])
        for i in range(2, 7):
            ws.append([f"2026-10-0{i}", f"Ürün {i}", i, 10 * i, f"=C{i}*D{i}"])
        ws.append(["Toplam", None, None, None, "=SUM(E2:E6)"])
    elif kind == "leave":
        from openpyxl.worksheet.datavalidation import DataValidation
        ws.title = "İzinler"
        ws.append(["Personel", "Başlangıç", "Bitiş", "Onaylandı"])
        dv = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("D2:D50")
    wb.save(out / name)
    _recalc(out / name)


# ---- PDF ------------------------------------------------------------------
def pdf_ops(up, out, file, ops):
    cur = str(up / file)
    dst = str(out / out_name(file, ".pdf"))
    for op in ops:
        k = op["op"]
        if k == "replace":
            F.replace_text(cur, dst, op["old"], op["new"], page=op.get("page"), near=op.get("near"), band=op.get("band"))
        elif k == "delete":
            F.delete_pages(cur, dst, op["pages"])
        elif k == "keep":
            F.keep_pages(cur, dst, op["pages"])
        elif k == "move":
            F.move_pages(cur, dst, op["pages"], op["before"])
        elif k == "insert":
            F.insert_pdf(cur, dst, str(up / op["other"]), at=op.get("at"))
        elif k == "rotate":
            F.rotate_pages(cur, dst, op["pages"], op["degrees"])
        elif k == "renumber":
            F.renumber(cur, dst)
        elif k == "watermark":
            _gotenberg_watermark(cur, dst, op["text"])
        elif k == "insert_paras":
            F.insert_paragraphs(cur, dst, op["after"], op["paragraphs"], page=op.get("page"))
        else:
            raise ValueError(k)
        cur = dst


def _gotenberg_watermark(src, dst, text):
    """Text watermark that keeps Turkish characters: the text is rendered to
    a one-page PDF by Chromium and applied with watermarkSource=pdf (the
    built-in text source drops glyphs such as 'İ')."""
    import requests
    base = os.environ.get("GOTENBERG_URL", "http://gotenberg:3000")
    html = ("<!doctype html><html><head><meta charset='utf-8'><style>@page{size:A4;margin:0} body{margin:0}"
            "div{margin-top:300px;text-align:center;font:bold 110px sans-serif;color:#C00000;opacity:.3;"
            f"transform:rotate(-35deg)}}</style></head><body><div>{text}</div></body></html>")
    wm = requests.post(base + "/forms/chromium/convert/html", files={"files": ("index.html", html.encode(), "text/html")},
                       timeout=120)
    wm.raise_for_status()
    with open(src, "rb") as fh:
        r = requests.post(base + "/forms/pdfengines/watermark",
                          files=[("files", (Path(src).name, fh, "application/pdf")), ("watermark", ("wm.pdf", wm.content, "application/pdf"))],
                          data={"watermarkSource": "pdf", "watermarkExpression": "wm.pdf"}, timeout=120)
    r.raise_for_status()
    Path(dst).write_bytes(r.content)


def pdf_create(up, out, name, title, color="#1F3864", table=False, landscape=False):
    import requests
    rows = "".join(f"<tr><td>Ürün {i}</td><td>{i * 10},00</td></tr>" for i in range(1, 6)) if table else ""
    html = (f"<!doctype html><html><head><meta charset='utf-8'><style>@page{{size:A4 {'landscape' if landscape else ''};}}"
            f"h1{{color:{color}}} table{{border-collapse:collapse}} td{{border:1px solid #888;padding:4px}}</style></head>"
            f"<body><h1>{title}</h1><p>Açıklama metni.</p>{'<table><tr><td>Ürün</td><td>Fiyat</td></tr>' + rows + '</table>' if table else ''}</body></html>")
    url = os.environ.get("GOTENBERG_URL", "http://gotenberg:3000") + "/forms/chromium/convert/html"
    r = requests.post(url, files={"files": ("index.html", html.encode(), "text/html")}, timeout=120)
    r.raise_for_status()
    (out / name).write_bytes(r.content)


# ---- misc -----------------------------------------------------------------
def copy(up, out, file):
    shutil.copyfile(up / file, out / file)


def none(up, out, **_):
    """Refusal: write nothing."""


def legacy(up, out, file, target, then):
    conv = up.parent / "scratch"
    conv.mkdir(exist_ok=True)
    subprocess.run([sys.executable, str(SCRIPTS / "convert_legacy.py"), str(up / file), "--outdir", str(conv)],
                   check=True, capture_output=True)
    converted = conv / (Path(file).stem + "." + target)
    tmp_up = up.parent / "legacy_up"
    tmp_up.mkdir(exist_ok=True)
    shutil.copyfile(converted, tmp_up / converted.name)
    globals()[then["fn"]](tmp_up, out, converted.name, **{k: v for k, v in then.items() if k not in ("fn", "file")})


def apply(ref, up, out):
    ref = dict(ref)
    fn = ref.pop("fn")
    globals()[fn](Path(up), Path(out), **ref)
