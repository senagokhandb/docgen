#!/usr/bin/env python3
"""DOCGEN recalc tests — xlsx_recalc.py with the IronCalc engine (docgen.calc).

1. Targeted cases on generated workbooks, checked against Excel's result:
   absolute relationship targets, external links, 3D references, AGGREGATE,
   TEXT rounding, unsupported functions, circular references, shared
   formulas with stale caches, tables, Turkish sheet names, value types,
   --expect, engine crash, package left byte-identical outside <v>.
2. Golden set: 269 formulas (tests/data/recalc_golden.json) — LibreOffice's
   results, replaced by Excel's where LibreOffice 7.4 differs.
3. With LibreOffice installed (agent image) and --fixtures: every fixture
   with formulas, inputs changed so caches are stale, recalculated by both
   engines; the results must be the same.

Usage:
    python tests/recalc_tests.py [--fixtures evals/fixtures] [--out DIR]
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "lib"))
from lxml import etree  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402
from openpyxl.worksheet.table import Table  # noqa: E402

from docgen import office  # noqa: E402
from docgen import xlsx as X  # noqa: E402
from docgen.package import Package, R_NS  # noqa: E402

S = X.s
RECALC = ROOT / "scripts" / "xlsx_recalc.py"
HAVE_LO = office.soffice() is not None


# ---- helpers --------------------------------------------------------------
def recalc(path, *args, engine="ironcalc"):
    r = subprocess.run([sys.executable, str(RECALC), str(path), "--json", "--engine", engine, *args],
                       capture_output=True, text=True)
    try:
        rep = json.loads(r.stdout)
    except ValueError:
        rep = {"raw": (r.stdout + r.stderr).strip()[-300:]}
    return r.returncode, rep


def cells(path):
    """{'Sheet!A1': (t, v)} of every formula cell as stored."""
    pkg = Package(path)
    sst = X.shared_strings(pkg)
    out = {}
    for name, part, _ in X.sheets(pkg):
        for c in pkg.xml(part).iter(S("c")):
            if c.find(S("f")) is None:
                continue
            t, v = c.get("t") or "n", c.findtext(S("v"))
            if t == "s":
                t, v = "str", sst[int(v)]
            out[f"{name}!{c.get('r')}"] = (t, v)
    return out


def same(got, want):
    """got: stored (t, v-text); want: (t, value)."""
    if got is None or got[0] != want[0]:
        return False
    if want[0] == "n":
        try:
            a, b = float(got[1]), float(want[1])
        except (TypeError, ValueError):
            return False
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    if want[0] == "b":
        return got[1] == ("1" if want[1] in (True, 1, "1") else "0")
    return (got[1] or "") == want[1]


def edit(path, fn):
    pkg = Package(path)
    fn(pkg)
    pkg.save(path)


def put_cache(pkg, target, value, t=None):
    """Write a stale cached value into a formula cell."""
    sheet, refs = X.parse_target(target)
    c = X._cell(pkg, sheet, refs[0])
    for v in c.findall(S("v")):
        c.remove(v)
    c.attrib.pop("t", None)
    if t:
        c.set("t", t)
    etree.SubElement(c, S("v")).text = str(value)
    pkg.touch(X.sheet_part(pkg, sheet))


def share(pkg, sheet, col, first, last, si):
    """Turn col{first}:col{last} into one shared formula group."""
    for r in range(first, last + 1):
        f = X._cell(pkg, sheet, f"{col}{r}").find(S("f"))
        f.set("t", "shared")
        f.set("si", str(si))
        if r == first:
            f.set("ref", f"{col}{first}:{col}{last}")
        else:
            f.text = None
    pkg.touch(X.sheet_part(pkg, sheet))


def add_external_link(pkg, book_target, sheet, cached):
    """[1]<sheet> external link whose cache holds {ref: number}."""
    main = X.S_NS
    part = "xl/externalLinks/externalLink1.xml"
    el = etree.Element(S("externalLink"), nsmap={None: main, "r": R_NS})
    book = etree.SubElement(el, S("externalBook"), {f"{{{R_NS}}}id": "rId1"})
    etree.SubElement(etree.SubElement(book, S("sheetNames")), S("sheetName"), val=sheet)
    sd = etree.SubElement(etree.SubElement(book, S("sheetDataSet")), S("sheetData"), sheetId="0")
    for ref, val in cached.items():
        row = etree.SubElement(sd, S("row"), r=re.sub(r"[A-Z]+", "", ref))
        etree.SubElement(etree.SubElement(row, S("cell"), r=ref), S("v")).text = str(val)
    pkg.write(part, el)
    pkg.set_override(part, "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml")
    pkg.add_rel(part, R_NS + "/externalLinkPath", book_target)
    rels = pkg.xml(pkg.rels_name(part))
    rels[0].set("TargetMode", "External")
    pkg.touch(pkg.rels_name(part))
    rid = pkg.add_rel(X.WB, R_NS + "/externalLink", "externalLinks/externalLink1.xml")
    wb = pkg.xml(X.WB)
    er = etree.Element(S("externalReferences"))
    etree.SubElement(er, S("externalReference"), {f"{{{R_NS}}}id": rid})
    wb.find(S("sheets")).addnext(er)
    pkg.touch(X.WB)


# ---- workbooks -----------------------------------------------------------
def book_main(path):
    """Shared formulas with stale caches, a table, Turkish names, comments
    (absolute rel targets), 3D references, an external link."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Satışlar"
    ws.append(["Bölge", "Ürün", "Adet", "Fiyat", "KDV", "Toplam", "Tutar"])
    data = [("Kuzey", "Kalem", 10, 2.5), ("Güney", "Defter", 5, 12), ("Kuzey", "Silgi", 20, 1.25), ("Doğu", "Kalem", 7, 2.5)]
    for i, (b, u, a, f) in enumerate(data, start=2):
        ws.append([b, u, a, f, f"=C{i}*D{i}*0.18", f"=C{i}*D{i}+E{i}", "=Satis[[#This Row],[Adet]]*Satis[[#This Row],[Fiyat]]"])
    ws.add_table(Table(displayName="Satis", ref="A1:G5"))
    ws["C2"].comment = Comment("girdi", "docgen")
    for n, v in (("Ay1", 10), ("Ay2", 20), ("Ay3", 30), ("Ay 4", 40), ("Ay 5", 50)):
        wb.create_sheet(n)["B2"] = v
    oz = wb.create_sheet("Özet")
    formulas = {
        "B1": "=SUM(Satışlar!F2:F5)", "B2": '=SUMIF(Satışlar!A2:A5,"Kuzey",Satışlar!F2:F5)', "B3": "=SUM(Satis[Tutar])",
        "B4": "=SUM(Ay1:Ay3!B2)", "B5": "=AVERAGE('Ay 4:Ay 5'!B2)", "B6": "=COUNT(Ay1:Ay3!B2)+B4",
        "B7": "=[1]Hedef!$B$2", "B8": "=[1]Hedef!$B$2*Parametre", "B9": "=B7-B1",
        "B10": '=TEXT(SUM(Satis[Adet]),"#,##0")', "B11": "=Satışlar!F2>100", "B12": '="Toplam: "&TEXT(B1,"#,##0.00")',
    }
    for ref, f in formulas.items():
        oz[ref] = f
    oz["D1"] = 2
    wb.defined_names["Parametre"] = DefinedName("Parametre", attr_text="'Özet'!$D$1")
    wb.save(path)

    def post(pkg):
        share(pkg, "Satışlar", "E", 2, 5, 0)
        share(pkg, "Satışlar", "F", 2, 5, 1)
        for r in range(2, 6):
            put_cache(pkg, f"Satışlar!E{r}", 1)          # stale on purpose
            put_cache(pkg, f"Satışlar!F{r}", 1)
        put_cache(pkg, "Özet!B7", 1500000)
        add_external_link(pkg, "Hedefler_2025.xlsx", "Hedef", {"B2": 1500000})
    edit(path, post)
    return {  # Excel's results
        "Satışlar!E2": ("n", 4.5), "Satışlar!F2": ("n", 29.5), "Satışlar!F3": ("n", 70.8), "Satışlar!F5": ("n", 20.65),
        "Satışlar!G4": ("n", 25), "Özet!B1": ("n", 150.45), "Özet!B2": ("n", 59.0), "Özet!B3": ("n", 127.5),
        "Özet!B4": ("n", 60), "Özet!B5": ("n", 45), "Özet!B6": ("n", 63), "Özet!B7": ("n", 1500000),
        "Özet!B8": ("n", 3000000), "Özet!B9": ("n", 1500000 - 150.45), "Özet!B10": ("str", "42"),
        "Özet!B11": ("b", False), "Özet!B12": ("str", "Toplam: 150.45"),
    }


def book_functions(path):
    """AGGREGATE, TEXT ties, Excel rounding, TODAY, types."""
    wb = Workbook()
    ws = wb.active
    ws.title = "F"
    for r, v in enumerate([10, 20, "=1/0", 30], start=1):
        ws.cell(r, 1).value = v
    cases = {
        "B1": ('=_xlfn.AGGREGATE(9,6,A1:A4)', ("n", 60)), "B2": ("=B1*2", ("n", 120)),
        "B3": ('=TEXT(12.5,"0")', ("str", "13")), "B4": ('=TEXT(2.675,"0.00")', ("str", "2.68")),
        "B5": ('=TEXT(0.125,"0%")', ("str", "13%")), "B6": ('=TEXT(1234.5,"#,##0")', ("str", "1,235")),
        "B7": ('=TEXT(DATE(2026,3,5),"dd.mm.yyyy")', ("str", "05.03.2026")), "B8": ('="Tutar: "&TEXT(12.5,"0")', ("str", "Tutar: 13")),
        "B9": ("=VALUE(B3)+1", ("n", 14)), "B10": ("=ROUND(48460*1.025,0)", ("n", 49672)),
        "B11": ("=ROUND(2.675,2)", ("n", 2.68)), "B12": ("=TODAY()", None), "B13": ("=1=1", ("b", True)),
        "B14": ('="x"&""', ("str", "x")), "B15": ('=IF(1>2,"a","")', ("str", "")), "B16": ("=Z99", ("n", 0)),
        "B17": ("=A3", ("e", "#DIV/0!")), "B18": ('=_xlfn.XLOOKUP(20,A1:A4,A1:A4)', ("n", 20)),
        "B19": ("=1/3", ("n", 1 / 3)), "B20": ('=UPPER("ışık istanbul")', ("str", "IŞIK ISTANBUL")),
        "B21": ('=_xlfn.NUMBERVALUE("12,5",",")', ("n", 12.5)), "B22": ('=_xlfn.NUMBERVALUE("1 234,5 ",","," ")', ("n", 1234.5)),
        "B23": ('="a"&TRUE', ("str", "aTRUE")), "B24": ("=TRUE>1", ("b", True)),
    }
    for ref, (f, _) in cases.items():
        ws[ref] = f
    wb.save(path)
    today = (dt.datetime.now(dt.timezone(dt.timedelta(hours=3))).date() - dt.date(1899, 12, 30)).days
    exp = {f"F!{k}": v for k, (_, v) in cases.items() if v is not None}
    exp["F!B12"] = ("n", today)
    exp["F!A3"] = ("e", "#DIV/0!")
    return exp


def book_unsupported(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "U"
    ws["A1"] = '=_xlfn.WEBSERVICE("https://example.com")'
    ws["A2"] = "=LEN(A1)"
    ws["A3"] = '=IFERROR(GETPIVOTDATA("Tutar",C1),0)'
    ws["A4"] = "=A3+1"
    wb.save(path)

    def post(pkg):
        put_cache(pkg, "U!A1", "eski", "str")
        put_cache(pkg, "U!A3", 7)
    edit(path, post)


def book_lookup(path):
    """LOOKUP over a computed array — the 'last match' idiom."""
    wb = Workbook()
    ws = wb.active
    ws.title = "L"
    for r, (b, n) in enumerate([("Kuzey", 10), ("Güney", 5), ("Kuzey", 8), ("Doğu", 20), ("Kuzey", 12), ("Güney", 7)], start=1):
        ws.cell(r, 1).value, ws.cell(r, 2).value = b, n
    ws["D1"] = '=LOOKUP(2,1/(A1:A6="Kuzey"),B1:B6)'
    ws["D2"] = "=D1*2"
    ws["D3"] = "=LOOKUP(9,B1:B6)"
    wb.save(path)
    edit(path, lambda p: put_cache(p, "L!D1", 99))


def book_circular(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "C"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "=B1+1", "=A1+1", "=5*2", "=FOOBAR(1)"
    wb.save(path)
    edit(path, lambda p: (put_cache(p, "C!A1", 1), put_cache(p, "C!B1", 2)))


def book_golden(path, golden):
    wb = Workbook()
    v = wb.active
    v.title = "Veri"
    v.append(["Tarih", "Bölge", "Ürün", "Adet", "Fiyat", "Tutar"])
    rows = [(dt.date(2026, 1, 5), "Kuzey", "Kalem", 10, 2.5), (dt.date(2026, 1, 12), "Güney", "Defter", 5, 12),
            (dt.date(2026, 2, 3), "Kuzey", "Defter", 8, 12), (dt.date(2026, 2, 17), "Doğu", "Silgi", 20, 1.25),
            (dt.date(2026, 3, 1), "Batı", "Kalem", 15, 2.5), (dt.date(2026, 3, 15), "Güney", "Kalem", 7, 2.5),
            (dt.date(2026, 4, 2), "Kuzey", "Silgi", 12, 1.25), (dt.date(2026, 4, 20), "Doğu", "Defter", 3, 12),
            (dt.date(2026, 5, 5), "Batı", "Defter", 9, 12), (dt.date(2026, 5, 25), "Güney", "Silgi", 30, 1.25)]
    for i, r in enumerate(rows, start=2):
        v.append(list(r))
        v.cell(i, 1).number_format = "dd.mm.yyyy"
    v.add_table(Table(displayName="Satis", ref="A1:F11"))
    v["H1"], v["I1"], v["H2"], v["H3"], v["H4"] = "KDV", 0.2, True, "42", 5
    wb.defined_names["KDV"] = DefinedName("KDV", attr_text="Veri!$I$1")
    wb.defined_names["Adetler"] = DefinedName("Adetler", attr_text="Veri!$D$2:$D$11")
    f = wb.create_sheet("Fiyatlar")
    f.append(["Ürün", "Fiyat", "Kategori", None, "Eşik", "Not"])
    for r in [("Kalem", 2.5, "Yazı", None, 0, "F"), ("Defter", 12, "Kağıt", None, 50, "C"),
              ("Silgi", 1.25, "Yazı", None, 70, "B"), (None, None, None, None, 85, "A")]:
        f.append(list(r))
    for i, (a, d) in enumerate([(-1000, dt.date(2026, 1, 1)), (300, dt.date(2026, 4, 1)),
                                (400, dt.date(2026, 9, 1)), (500, dt.date(2027, 3, 1))], start=1):
        f.cell(i, 8).value, f.cell(i, 9).value = a, d
    e = wb.create_sheet("Ek Sayfa")
    e["A1"], e["A2"] = 100, 200
    s = wb.create_sheet("Satışlar")
    s["B2"], s["B3"], s["B4"] = 1.5, 2.5, 3
    for n, val in (("Ay1", 10), ("Ay2", 20), ("Ay3", 30)):
        wb.create_sheet(n)["B2"] = val
    wb.create_sheet("Test")
    wb.create_sheet("Ozet")
    for g in golden:
        sheet, ref = g["cell"].split("!")
        wb[sheet][ref] = ArrayFormula(ref, g["formula"]) if g["kind"] == "array" else g["formula"]
    wb.save(path)


def book_big(path, n=20000):
    wb = Workbook()
    ws = wb.active
    ws.title = "Satışlar"
    for i in range(2, n + 2):
        ws.append(["Kuzey Güney Doğu Batı Merkez".split()[i % 5], (i * 7) % 50 + 1, ((i * 13) % 997) / 10 + 1,
                   f"=B{i}*C{i}", f"=ROUND(D{i}*1.2,2)", f'=IF(E{i}>1000,"Büyük","Küçük")'])
    oz = wb.create_sheet("Özet")
    for j, b in enumerate("Kuzey Güney Doğu Batı Merkez".split(), start=1):
        oz.cell(j, 1).value = b
        oz.cell(j, 2).value = f"=SUMIF(Satışlar!$A$1:$A${n + 1},A{j},Satışlar!$E$1:$E${n + 1})"
    wb.save(path)


def broken_styles(path):
    """A package IronCalc cannot load (styles without cellStyles)."""
    wb = Workbook()
    wb.active["A1"], wb.active["A2"] = 2, "=A1*3"
    wb.save(path)

    def post(pkg):
        st = pkg.xml("xl/styles.xml")
        for el in list(st):
            if etree.QName(el).localname in ("cellStyles", "cellXfs", "cellStyleXfs"):
                st.remove(el)
        pkg.touch("xl/styles.xml")
    edit(path, post)


def stale_copy(src, dst):
    """Double every numeric constant so every formula cache is stale."""
    pkg = Package(src)
    for _, part, _ in X.sheets(pkg):
        root, hit = pkg.xml(part), False
        for c in root.iter(S("c")):
            v = c.find(S("v"))
            if c.find(S("f")) is None and c.get("t") in (None, "n") and v is not None and v.text:
                num = float(v.text) * 2
                v.text = str(int(num)) if num.is_integer() else repr(num)
                hit = True
        if hit:
            pkg.touch(part)
    pkg.save(dst)


# ---- tests ---------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--fixtures", help="dir with the eval fixtures (LibreOffice comparison)")
    ap.add_argument("--out", help="keep the test workbooks here")
    a = ap.parse_args()
    out = Path(a.out or tempfile.mkdtemp(prefix="recalc-tests-"))
    out.mkdir(parents=True, exist_ok=True)
    rows = []

    def report(tid, ok, detail=""):
        rows.append((tid, "OK" if ok else "FAIL", detail))

    def check_values(tid, path, expected):
        got = cells(path)
        bad = [f"{k}: {got.get(k)} != {w}" for k, w in expected.items() if not same(got.get(k), w)]
        report(tid, not bad, f"{len(expected) - len(bad)}/{len(expected)} as Excel" + (f"; {bad[:3]}" if bad else ""))

    # 1. main workbook: rels, external link, 3D, tables, shared formulas, Turkish names
    p = out / "main.xlsx"
    expected = book_main(p)
    before = {i.filename: zipfile.ZipFile(p).read(i.filename) for i in zipfile.ZipFile(p).infolist()}
    rc, rep = recalc(p)
    report("load: comments with absolute rel targets", rc == 0 and rep.get("engine") == "ironcalc", f"rc={rc} {rep.get('engine')}")
    check_values("values: shared/table/3D/external/TEXT", p, expected)
    after = {i.filename: zipfile.ZipFile(p).read(i.filename) for i in zipfile.ZipFile(p).infolist()}
    sheet_parts = {part for _, part, _ in X.sheets(Package(p))}
    leaked = sorted(set(after) ^ set(before))
    changed = sorted(n for n in before if n in after and before[n] != after[n] and n not in sheet_parts)
    report("package: only sheet <v> changed", not leaked and not changed, f"new/removed={leaked} changed={changed}")
    rc2, rep2 = recalc(p, "--expect", "Özet!B4=60", "--expect", "Özet!B1=999")
    report("--expect", rc2 == 1 and len(rep2.get("failed_expectations", [])) == 1, str(rep2.get("failed_expectations")))

    # agent path: formulas written by docgen.xlsx (no cached value), an input changed
    p = out / "agent.xlsx"
    wb = Workbook()
    wb.active.title = "Notlar"
    for r, (v, f, o) in enumerate([(92, 94, 48), (52, 63, 43), (46, 38, 55)], start=2):
        wb.active.append([None]) if r == 2 else None
        wb.active.cell(r, 3).value, wb.active.cell(r, 4).value, wb.active.cell(r, 5).value = v, f, o
    wb.save(p)
    pkg = Package(p)
    for r in range(2, 5):
        X.set_formula(pkg, f"Notlar!F{r}", f"=ROUND(C{r}*0.3+D{r}*0.5+E{r}*0.2,1)")
        X.set_formula(pkg, f"Notlar!G{r}", f'=IF(F{r}>=50,"Geçti","Kaldı")')
    X.set_value(pkg, "Notlar!C4", 60, expect=46)
    pkg.save(p)
    rc, rep = recalc(p)
    got = cells(p)
    want = {"Notlar!F2": ("n", 84.2), "Notlar!F3": ("n", 55.7), "Notlar!F4": ("n", 48), "Notlar!G4": ("str", "Kaldı"),
            "Notlar!G2": ("str", "Geçti")}
    bad = [f"{k}: {got.get(k)}" for k, w in want.items() if not same(got.get(k), w)]
    report("agent path: set_formula/set_value", rc == 0 and rep.get("engine") == "ironcalc" and not bad,
           f"rc={rc} {rep.get('engine') or rep.get('raw', '')[:100]} {bad}")

    # 2. functions: AGGREGATE/TEXT via Formualizer, Excel rounding, types
    p = out / "functions.xlsx"
    expected = book_functions(p)
    rc, rep = recalc(p)
    check_values("values: AGGREGATE/TEXT/ROUND/types", p, expected)
    report("new error reported (=1/0)", rc == 1 and any("F!A3" in e for e in rep.get("new_errors", [])), str(rep.get("new_errors")))

    # 3. unsupported functions: cached value kept, dependents use it, FAIL
    p = out / "unsupported.xlsx"
    book_unsupported(p)
    rc, rep = recalc(p)
    nr = rep.get("not_recalculated", [])
    got = cells(p)
    report("unsupported: reported, cache kept", rc == 1 and len(nr) == 2 and got["U!A1"] == ("str", "eski")
           and same(got["U!A3"], ("n", 7)), str(nr))
    report("unsupported: dependents use the kept value", same(got["U!A2"], ("n", 4)) and same(got["U!A4"], ("n", 8)),
           f"A2={got['U!A2']} A4={got['U!A4']}")

    # LOOKUP over a computed array: LibreOffice when installed, otherwise reported
    p = out / "lookup.xlsx"
    book_lookup(p)
    rc, rep = recalc(p)
    got = cells(p)
    if HAVE_LO:
        report("LOOKUP(2,1/(...)) via LibreOffice", rc == 0 and same(got["L!D1"], ("n", 12)) and same(got["L!D2"], ("n", 24))
               and same(got["L!D3"], ("n", 8)), f"D1={got['L!D1']} D2={got['L!D2']} {rep.get('notes')}")
    else:
        report("LOOKUP(2,1/(...)) reported, cache kept", rc == 1 and len(rep.get("not_recalculated", [])) == 1
               and same(got["L!D1"], ("n", 99)) and same(got["L!D2"], ("n", 198)) and same(got["L!D3"], ("n", 8)),
               f"{rep.get('not_recalculated')} D2={got['L!D2']}")

    # 4. circular reference and an unknown function
    p = out / "circular.xlsx"
    book_circular(p)
    rc, rep = recalc(p)
    got = cells(p)
    report("circular: reported, cache kept", rc == 1 and got["C!A1"] == ("n", "1")
           and any("circular" in x for x in rep.get("not_recalculated", [])) and same(got["C!C1"], ("n", 10)),
           str(rep.get("not_recalculated")))
    report("unknown function -> #NAME? (as Excel)", got["C!D1"] == ("e", "#NAME?"), str(got["C!D1"]))

    # 5. engine cannot open the file: LibreOffice fallback, or a clean FAIL
    p = out / "broken.xlsx"
    broken_styles(p)
    rc, rep = recalc(p)
    if HAVE_LO:
        report("engine crash -> LibreOffice fallback", rc == 0 and rep.get("engine") == "libreoffice", str(rep)[:120])
    else:
        report("engine crash -> clean FAIL", rc == 1 and "raw" in rep and "FAIL" in rep["raw"], str(rep)[:120])

    # 6. .xlsm
    p = out / "macro.xlsm"
    book_functions(out / "m.xlsx")
    zin, zout = zipfile.ZipFile(out / "m.xlsx"), zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED)
    for i in zin.infolist():
        d = zin.read(i.filename)
        zout.writestr(i, d.replace(b"sheet.main+xml", b"sheet.macroEnabled.main+xml") if i.filename == "[Content_Types].xml" else d)
    zout.close()
    rc, rep = recalc(p)
    report(".xlsm", rep.get("engine") == "ironcalc" and same(cells(p).get("F!B1"), ("n", 60)), str(rep.get("engine")))

    # 7. golden set
    golden = json.loads((ROOT / "tests" / "data" / "recalc_golden.json").read_text())
    p = out / "golden.xlsx"
    book_golden(p, golden)
    recalc(p)
    got = cells(p)
    bad = [f"{g['cell']} {g['formula'][:40]}: {got.get(g['cell'])} != {g['expect']}" for g in golden
           if not same(got.get(g["cell"]), tuple(g["expect"]))]
    report(f"golden set ({len(golden)} formulas)", not bad, f"{len(golden) - len(bad)}/{len(golden)}" + (f"; {bad[:4]}" if bad else ""))

    # 8. size
    p = out / "big.xlsx"
    book_big(p)
    t0 = time.time()
    rc, rep = recalc(p)
    secs = time.time() - t0
    report("60k formulas", rc == 0 and secs < 30, f"{secs:.1f}s, {rep.get('formulas')} results written")

    # 9. same results as LibreOffice on the fixtures, caches stale
    if HAVE_LO and a.fixtures:
        for fx in sorted(Path(a.fixtures).glob("*.xlsx")):
            if not any(b"<f" in zipfile.ZipFile(fx).read(n) for n in zipfile.ZipFile(fx).namelist()
                       if n.startswith("xl/worksheets/")):
                continue
            ic_p, lo_p = out / f"fx_ic_{fx.name}", out / f"fx_lo_{fx.name}"
            stale_copy(fx, ic_p)
            shutil.copy(ic_p, lo_p)
            rc1, r1 = recalc(ic_p)
            rc2, r2 = recalc(lo_p, engine="libreoffice")
            a_, b_ = cells(ic_p), cells(lo_p)
            diff = [f"{k}: {a_[k]} vs LO {b_.get(k)}" for k in a_
                    if not (a_[k] == b_.get(k) or (a_[k][0] == "n" and b_.get(k, ("",))[0] == "n"
                                                    and same(a_[k], ("n", float(b_[k][1])))))]
            report(f"= LibreOffice: {fx.name}", not diff and rc1 == rc2, f"{len(a_)} formulas, rc {rc1}/{rc2}"
                   + (f"; {diff[:3]}" if diff else ""))
    elif a.fixtures:
        rows.append(("= LibreOffice: fixtures", "SKIPPED", "LibreOffice not installed"))

    w = max(len(r[0]) for r in rows)
    for tid, status, detail in rows:
        print(f"{tid:<{w}}  {status:<7}  {detail}")
    bad = sum(1 for r in rows if r[1] == "FAIL")
    print(f"\n{len(rows) - bad}/{len(rows)} OK  (workbooks in {out})")
    if not a.out:
        shutil.rmtree(out, ignore_errors=True)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
